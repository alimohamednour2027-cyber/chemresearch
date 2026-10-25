import streamlit as st
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

st.set_page_config(
    page_title="Clean Lab Graphical Abstract",
    page_icon="🧪",
    layout="wide"
)

st.title("🧪 ChemResearch AI: Professional Graphical Abstract Generator")
st.markdown("هذا المحرك يولد مخططاً أفقياً واسعاً ونظيفاً (يعتمد على الكلمات المفتاحية والأيقونات البصرية) لتجنب أي تداخل أو نصوص طويلة، ومطابق لاشتراطات المجلات العلمية.")
st.markdown("---")

st.subheader("✍️ أدخل خطوات العملية المخبرية (مختصرة ومركزة لكل مرحلة):")
default_steps = """1. Stage 1: Raw Biomass Washing
2. Stage 2: Oven Drying @ 100°C
3. Stage 3: Chemical Treatment & Stirring
4. Stage 4: Final Pure Bio-Adsorbent Yield"""

steps_input = st.text_area("خطوات البروتوكول المخبري:", value=default_steps, height=160)
abstract_title = st.text_input("عنوان المخطط الرسومي:", value="Extraction and Synthesis Workflow of Bio-Adsorbent")

if st.button("🚀 توليد وتحميل ملف الـ PowerPoint النظيف", key="gen_clean_ppt"):
    if not steps_input.strip():
        st.error("الرجاء إدخال خطوة واحدة على الأقل.")
    else:
        with st.spinner("جاري تصميم الشريحة الهندسية النظيفة..."):
            prs = Presentation()
            prs.slide_width = Inches(13.33)  # قياس عرض الشريحة 16:9
            prs.slide_height = Inches(7.5)
            
            blank_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_layout)
            
            # عنوان الشريحة الأكاديمي
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.73), Inches(0.8))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            p_title = tf_title.paragraphs[0]
            p_title.text = abstract_title
            p_title.font.name = "Times New Roman"
            p_title.font.size = Pt(22)
            p_title.font.bold = True
            p_title.font.color.rgb = RGBColor(20, 20, 20)
            
            steps = [line.strip() for line in steps_input.split('\n') if line.strip()]
            num_steps = len(steps)
            
            if num_steps > 0:
                start_left = Inches(0.8)
                start_top = Inches(1.8)
                total_width = Inches(11.73)
                
                box_width = (total_width / num_steps) - Inches(0.25)
                box_height = Inches(4.8)
                spacing = Inches(0.25)
                
                for i, step_text in enumerate(steps):
                    current_left = start_left + i * (box_width + spacing)
                    
                    # 1. صندوق الخلفية البيضاء النظيفة لكل مرحلة
                    card = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE, 
                        current_left, start_top, box_width, box_height
                    )
                    card.fill.solid()
                    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    card.line.color.rgb = RGBColor(180, 190, 200)
                    card.line.width = Pt(1.2)
                    
                    # 2. مربع مخصص للأيقونة البصرية في أعلى الصندوق
                    icon_size = Inches(1.4)
                    icon_left = current_left + (box_width / 2) - (icon_size / 2)
                    icon_top = start_top + Inches(0.3)
                    
                    icon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, icon_left, icon_top, icon_size, icon_size)
                    icon.fill.solid()
                    icon.fill.fore_color.rgb = RGBColor(235, 242, 250)
                    icon.line.color.rgb = RGBColor(150, 170, 190)
                    
                    # 3. صندوق النص المختصر والمركز أسفل الأيقونة
                    tb = slide.shapes.add_textbox(
                        current_left + Inches(0.1), 
                        start_top + Inches(1.8), 
                        box_width - Inches(0.2), 
                        Inches(2.8)
                    )
                    tf = tb.text_frame
                    tf.word_wrap = True
                    
                    # تقسيم النص إلى عنوان المرحلة ووصفها المختصر إذا وجد نقطتان
                    if ":" in step_text:
                        parts = step_text.split(":", 1)
                        stage_label = parts[0].strip()
                        stage_desc = parts[1].strip()
                    else:
                        stage_label = f"Stage {i+1}"
                        stage_desc = step_text
                    
                    # عنوان المرحلة (مثل Stage 1)
                    p_num = tf.paragraphs[0]
                    p_num.text = stage_label
                    p_num.font.name = "Times New Roman"
                    p_num.font.size = Pt(13)
                    p_num.font.bold = True
                    p_num.font.color.rgb = RGBColor(50, 90, 140)
                    p_num.alignment = PP_ALIGN.CENTER
                    
                    # الوصف العلمي المختصر والموجز
                    p_desc = tf.add_paragraph()
                    p_desc.text = stage_desc
                    p_desc.font.name = "Times New Roman"
                    p_desc.font.size = Pt(11)
                    p_desc.font.color.rgb = RGBColor(40, 40, 40)
                    p_desc.alignment = PP_ALIGN.CENTER
                    
                    # 4. سهم التتابع الهندسي بمسافة آمنة وواضحة
                    if i < num_steps - 1:
                        arrow_left = current_left + box_width + Inches(0.04)
                        arrow_top = start_top + (box_height / 2) - Inches(0.2)
                        arrow = slide.shapes.add_shape(
                            MSO_SHAPE.RIGHT_ARROW,
                            arrow_left, arrow_top, Inches(0.18), Inches(0.18)
                        )
                        arrow.fill.solid()
                        arrow.fill.fore_color.rgb = RGBColor(110, 110, 110)
                        arrow.line.color.rgb = RGBColor(80, 80, 80)

            # حفظ الملف في الذاكرة لتنزيله فوراً
            buffer = BytesIO()
            prs.save(buffer)
            buffer.seek(0)
            
            st.success("تم توليد ملف الـ PowerPoint النظيف بنجاح!")
            st.download_button(
                label="📥 تنزيل ملف PowerPoint النموذجي (.pptx)",
                data=buffer,
                file_name="Clean_Graphical_Abstract_Example.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
