import streamlit as st
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pypdf
from google import genai

# إعدادات صفحة الموقع
st.set_page_config(
    page_title="ChemResearch AI Platform",
    layout="wide"
)

# --- تنسيق CSS المخصص للهيدر العلوي والشريط الجانبي (بدون إيقونات، أزرق داكن وفاتح) ---
st.markdown("""
    <style>
    /* الهيدر العلوي - أزرق داكن مع خط أبيض غليظ */
    .stApp > header {
        background-color: transparent !important;
    }
    .custom-header {
        background-color: #1a365d;
        color: #ffffff;
        padding: 25px;
        border-radius: 8px;
        margin-bottom: 20px;
    }
    .custom-header h1 {
        color: #ffffff;
        font-weight: 800;
        margin: 0;
        font-size: 28px;
    }
    .custom-header p {
        color: #e2e8f0;
        margin: 5px 0 0 0;
        font-size: 15px;
    }
    
    /* الشريط الجانبي - أزرق أقل حدة */
    [data-testid="stSidebar"] {
        background-color: #2b4c7e;
        color: #ffffff;
    }
    [data-testid="stSidebar"] .stMarkdown, [data-testid="stSidebar"] label, [data-testid="stSidebar"] span {
        color: #ffffff !important;
    }
    </style>
""", unsafe_allow_html=True)

# عرض الهيدر الاحترافي
st.markdown("""
    <div class="custom-header">
        <h1>ChemResearch AI Platform</h1>
        <p>Advanced Academic Platform & Professional Q1 Research Synthesis Engine</p>
    </div>
""", unsafe_allow_html=True)

# ==================== الحفاظ على البيانات عند تغيير الخيارات (Session State) ====================
if "intro_output" not in st.session_state:
    st.session_state.intro_output = ""

# ==================== الشريط الجانبي وإعدادات الـ API ====================
with st.sidebar:
    st.markdown("### Language Configuration")
    lang_choice = st.selectbox("Choose Language:", ["English", "Türkçe", "العربية"], label_visibility="collapsed")

    st.markdown("---")
    st.markdown("### Gemini API Configuration")
    api_key_input = st.text_input("Enter your Gemini API Key:", type="password", value="")

    st.markdown("---")
    st.markdown("### Literature Hub")
    t_upload_sec2 = "Upload Review Papers (PDFs)"

# قاموس الترجمات الشامل (خالٍ تماماً من الإيقونات)
if lang_choice == "English":
    sec1_title = "Section 1: Experiment & Research Metadata"
    sec2_title = "Section 2: AI-Synthesized Introduction & Professional Citations"
    t_paper_title = "Article or Thesis Title:"
    t_target_metal = "Target Heavy Metal / Pollutant:"
    t_adsorbent = "Prepared Novel Material / Sorbent:"
    t_tools = "Characterization Tools (e.g., FTIR, SEM):"
    t_analytical_inst = "Analytical Instrumentation (e.g., AAS, UV-Vis):"
    t_fields_label = "Select Research Fields:"
    t_field_1 = "Solid-Phase Extraction (SPE)"
    t_field_2 = "Adsorption"
    t_field_3 = "Removal / Remediation"
    t_field_other = "Other (Specify):"
    t_q_style = "Target Journal Quartile Style:"
    t_citation_style = "Citation Style:"
    t_btn_sec2 = "Synthesize Introduction & Professional Citations"
    t_spinner = "AI is analyzing your PDFs and generating a professional academic synthesis..."
    t_success_gen = "Introduction successfully synthesized with professional citations."
    t_download_btn = "Download Introduction & References as Word (.docx)"

elif lang_choice == "Türkçe":
    sec1_title = "Bolum 1: Deney ve Arastirma Verileri"
    sec2_title = "Bolum 2: Yapay Zeka Destekli Giris ve Profesyonel Atif Merkezi"
    t_paper_title = "Makale veya Tez Basligi:"
    t_target_metal = "Hedef Agir Metal / Kirletici:"
    t_adsorbent = "Hazirlanan Yeni Malzeme / Adsorban:"
    t_tools = "Karakterizasyon Cihazlari (orn. FTIR, SEM):"
    t_analytical_inst = "Analitik Olcum Cihazi (orn. AAS, UV-Vis):"
    t_fields_label = "Arastirma Alanlarini Secin:"
    t_field_1 = "Kati Faz Ekstraksiyonu (SPE)"
    t_field_2 = "Adsorpsiyon"
    t_field_3 = "Giderim / Aritma"
    t_field_other = "Diger (Belirtiniz):"
    t_q_style = "Hedef DergiDuzeyi (Quartile):"
    t_citation_style = "Atif Stili:"
    t_btn_sec2 = "Giris ve Profesyonel Atiflari Sentezle"
    t_spinner = "Yapay zeka PDF'lerinizi analiz ediyor ve profesyonel bir sentez olusturuyor..."
    t_success_gen = "Giris ve profesyonel atiflar basariyla sentezlendi."
    t_download_btn = "Giris ve Referanslari Word Olarak Indir (.docx)"

else:  # العربية
    sec1_title = "القسم الأول: بيانات التجربة والبحث الأساسية"
    sec2_title = "القسم الثاني: مركز صياغة المقدمة والتوثيق العلمي الاحترافي"
    t_paper_title = "عنوان الأطروحة أو المقال العلمي:"
    t_target_metal = "المعدن الثقيل / الملوث المستهدف:"
    t_adsorbent = "المادة الماصة / الجديدة المُحضرة:"
    t_tools = "أجهزة التشخيص (مثل: FTIR, SEM, XRD):"
    t_analytical_inst = "جهاز التحليل المخبري لتقدير التراكيز (مثل: AAS, UV-Vis, ICP-OES):"
    t_fields_label = "اختر مجالات البحث الأساسية:"
    t_field_1 = "الاستخلاص الصلب (Solid-Phase Extraction - SPE)"
    t_field_2 = "الامتزاز (Adsorption)"
    t_field_3 = "الإزالة والمعالجة (Removal / Remediation)"
    t_field_other = "أخرى (اكتبها بنفسك):"
    t_q_style = "مستوى المجلة المستهدفة (Journal Quartile Style):"
    t_citation_style = "أسلوب التوثيق والاستشهاد (Citation Style):"
    t_btn_sec2 = "صياغة المقدمة والتوثيق العلمي الاحترافي"
    t_spinner = "الذكاء الاصطناعي يحلل مقالات المراجعة ويوثق النص بأسلوب المجلات العالمية..."
    t_success_gen = "تمت صياغة المقدمة وتوثيقها بالمراجع بأسلوب علمي احترافي ناجح."
    t_download_btn = "تحميل المقدمة وقائمة المراجع كملف Word نظيف (.docx)"

# ==================== القسم الأول: بيانات التجربة والبحث ====================
st.markdown(f"### {sec1_title}")

col1, col2 = st.columns(2)
with col1:
    paper_title = st.text_input(t_paper_title, value="Facile synthesis of a novel bio-adsorbent for cadmium removal from aqueous solutions")
    target_metal = st.text_input(t_target_metal, value="Cadmium(II) - Cd2+")
    analytical_instrument = st.text_input(t_analytical_inst, value="Atomic Absorption Spectrometer (AAS)")
with col2:
    adsorbent_material = st.text_input(t_adsorbent, value="Modified Agricultural Waste Bio-adsorbent")
    characterization_tools = st.text_input(t_tools, value="FTIR, SEM, XRD, BET")

st.markdown(f"**{t_fields_label}**")
col_f1, col_f2, col_f3 = st.columns(3)
with col_f1: f_spe = st.checkbox(t_field_1, value=True)
with col_f2: f_ads = st.checkbox(t_field_2, value=True)
with col_f3: f_rem = st.checkbox(t_field_3, value=True)

other_field_input = st.text_input(t_field_other, value="")
journal_style = st.selectbox(t_q_style, ["Q1 Style (High Impact & Rigorous)", "Q2 Style (Standard Professional)", "Q3 Style (Clear & Direct)"])
citation_format = st.selectbox(t_citation_style, ["Numerical Style [1], [2] (Standard Journal Format)", "Author-Date Style (APA / Harvard)"])

selected_fields = []
if f_spe: selected_fields.append("Solid-Phase Extraction (SPE)")
if f_ads: selected_fields.append("Adsorption")
if f_rem: selected_fields.append("Removal/Remediation")
if other_field_input: selected_fields.append(other_field_input)
fields_text = ", ".join(selected_fields)

st.markdown("---")

# ==================== القسم الثاني: مركز صياغة المقدمة الذكي التلقائي ====================
st.markdown(f"### {sec2_title}")

st.markdown("مستودع مقالات المراجعة (ارفع ملفات الـ PDF هنا):")
sources_sec2 = st.file_uploader(t_upload_sec2, type=["pdf", "txt"], accept_multiple_files=True, key="sec2_files")

extracted_pdf_text = ""
if sources_sec2:
    for f in sources_sec2:
        try:
            reader = pypdf.PdfReader(f)
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    extracted_pdf_text += f"\n--- Source File: {f.name} (Page {page_idx+1}) ---\n" + page_text + "\n"
        except Exception as e:
            st.error(f"Error reading PDF: {e}")
    st.success(f"Successfully extracted text from {len(sources_sec2)} review files.")

if st.button(t_btn_sec2, key="btn_sec2"):
    if not api_key_input:
        st.error("Please enter your Gemini API Key in the sidebar to enable AI synthesis.")
    else:
        with st.spinner(t_spinner):
            try:
                client = genai.Client(api_key=api_key_input)
                
                prompt = f"""
                You are an expert analytical chemistry researcher and professional academic writer publishing in top Q1 journals.
                Your task is to write a rigorous, comprehensive, and entirely original Introduction section for a PhD-level research article.
                
                PROFESSIONAL CITATION & FORMATTING RULES:
                1. STRICTLY DO NOT use Markdown asterisks (*) or bold/italic markdown symbols anywhere in the text (no *, no **). Keep plain clean text.
                2. Use natural academic citation distribution: A reference can be cited where appropriate across different sections of the introduction.
                3. Build a logical funnel introduction (Global environmental impact -> Limitations of conventional methods -> Research Gap -> Novelty of {adsorbent_material}).
                4. At the end of the text, provide a clean References section listing unique literature derived from the uploaded sources in correct academic order, without any markdown stars.
                5. Avoid complex LaTeX formatting (use standard text formats like Cu2+, Cd2+, qm, R2 > 0.99).
                
                Research Metadata:
                - Title: {paper_title}
                - Target Pollutant: {target_metal}
                - Prepared Material/Sorbent: {adsorbent_material}
                - Fields: {fields_text}
                - Characterization Tools: {characterization_tools}
                - Journal Standard: {journal_style}
                - Citation Style: {citation_format}
                
                Extracted Text from Uploaded Review Papers:
                {extracted_pdf_text[:25000]}
                
                Write in formal academic English suitable for Q1 journals.
                """
                
                response = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=prompt,
                )
                
                cleaned_text = response.text.replace('*', '').replace('###', '')
                st.session_state.intro_output = cleaned_text
            except Exception as e:
                st.error(f"An error occurred during AI generation: {e}")

# عرض النتائج
if st.session_state.intro_output:
    st.markdown("#### Generated Introduction & Professional Citations Preview:")
    st.markdown(st.session_state.intro_output)
    st.success(t_success_gen)

    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1) if 'Inches' in globals() else doc.sections[0].top_margin
        section.bottom_margin = doc.sections[0].bottom_margin
        section.left_margin = doc.sections[0].left_margin
        section.right_margin = doc.sections[0].right_margin

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    run_title = p_title.add_run(paper_title + " - Introduction & References")
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)

    final_clean_text = st.session_state.intro_output.replace('*', '')
    paragraphs = final_clean_text.split('\n\n')
    
    for para_text in paragraphs:
        if para_text.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(para_text.replace('#', '').strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    st.download_button(
        label=t_download_btn,
        data=buffer,
        file_name="ChemResearch_AI_Professional_Introduction.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        key="dl_sec2"
    )
